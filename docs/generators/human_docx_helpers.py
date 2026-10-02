import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# PROFESSIONAL CORPORATE PALETTE (Executive Navy & Slate)
# ==============================================================================
COLOR_PRIMARY = RGBColor(15, 41, 66)      # Deep Midnight Navy (#0F2942)
COLOR_SECONDARY = RGBColor(30, 58, 138)   # Royal Navy (#1E3A8A)
COLOR_DARK = RGBColor(30, 41, 59)         # Dark Charcoal Slate (#1E293B)
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray (#64748B)
COLOR_CODE = RGBColor(159, 18, 57)        # Wine Red (#9F1239)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def add_total_pages(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'NUMPAGES'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def setup_page_layout(doc, doc_title):
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.different_first_page_header_footer = True
    
    # Running Header (Pages 2+)
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"Getthawha Wellness Platform  |  {doc_title}")
    hrun.font.name = 'Calibri'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED

    # Running Footer (Pages 2+)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun1 = fp.add_run("Commercial Confidential — For Internal & Evaluative Use Only\tPage ")
    frun1.font.name = 'Calibri'
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = COLOR_MUTED
    
    add_page_number(frun1)
    frun2 = fp.add_run(" of ")
    frun2.font.name = 'Calibri'
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = COLOR_MUTED
    add_total_pages(frun2)

def create_cover_page(doc, doc_title, subtitle, meta_items):
    # Top spacing
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_before = Pt(40)
    p_org.paragraph_format.space_after = Pt(8)
    r_org = p_org.add_run("GETTHAWA THAI MASSAGE & WELLNESS GROUP")
    r_org.font.name = 'Calibri'
    r_org.font.size = Pt(10.5)
    r_org.font.bold = True
    r_org.font.color.rgb = COLOR_SECONDARY

    # Thin accent rule under org name
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(24)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="0F2942"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run(doc_title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_MUTED

    # Metadata Card (Table)
    tbl = doc.add_table(rows=len(meta_items), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(tbl, color="E2E8F0", sz="4")
    
    col_widths = [Inches(2.2), Inches(4.3)]
    for r_idx, (k, v) in enumerate(meta_items):
        row = tbl.rows[r_idx]
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]
        set_cell_margins(row.cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(row.cells[1], top=80, bottom=80, left=120, right=120)
        set_cell_background(row.cells[0], "F8FAFC")
        set_cell_background(row.cells[1], "FFFFFF")

        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    # Push to next page
    doc.add_page_break()

def create_control_page(doc, revisions, approvals):
    add_heading(doc, "Document Control & Approvals", level=1)
    
    p = doc.add_paragraph("This document represents an official engineering deliverable for the Getthawha Platform. Formal sign-off and historical revisions are cataloged below:")
    p.paragraph_format.space_after = Pt(12)

    # Approvals Table
    add_heading(doc, "Formal Review & Acceptance Sign-Off", level=2)
    tbl_app = doc.add_table(rows=len(approvals)+1, cols=5)
    tbl_app.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_app)

    headers_app = ["Role / Representation", "Name", "Title / Responsibility", "Approval Status", "Sign Date"]
    for c_idx, h in enumerate(headers_app):
        cell = tbl_app.rows[0].cells[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, app in enumerate(approvals):
        row = tbl_app.rows[r_idx+1]
        for c_idx, val in enumerate(app):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(12)

    # Revision History
    add_heading(doc, "Document Revision History", level=2)
    tbl_rev = doc.add_table(rows=len(revisions)+1, cols=4)
    tbl_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rev)

    headers_rev = ["Version", "Release Date", "Author / Role", "Summary of Technical Changes"]
    for c_idx, h in enumerate(headers_rev):
        cell = tbl_rev.rows[0].cells[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, rev in enumerate(revisions):
        row = tbl_rev.rows[r_idx+1]
        for c_idx, val in enumerate(rev):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    doc.add_page_break()

def create_toc_page(doc, toc_items):
    add_heading(doc, "Table of Contents", level=1)
    p_lead = doc.add_paragraph("The structure of this document is organized into the following principal sections:")
    p_lead.paragraph_format.space_after = Pt(14)

    for num_str, title_str in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        r_num = p.add_run(f"{num_str}\t")
        r_num.font.name = 'Calibri'
        r_num.font.bold = True
        r_num.font.size = Pt(10.5)
        r_num.font.color.rgb = COLOR_PRIMARY

        r_title = p.add_run(title_str)
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = COLOR_DARK

    doc.add_page_break()

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        # Bottom border on H1
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="0F2942"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_MUTED

def add_p(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Calibri'
        r_pre.font.bold = True
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_DARK
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Calibri'
        r_pre.font.bold = True
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

def add_callout(doc, text, title="OPERATIONAL CONTEXT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="0F2942"/>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.15
    
    r_tit = cp.add_run(f"[{title}] ")
    r_tit.font.name = 'Calibri'
    r_tit.font.size = Pt(9.5)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    r_txt = cp.add_run(text)
    r_txt.font.name = 'Calibri'
    r_txt.font.size = Pt(9.5)
    r_txt.font.italic = True
    r_txt.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)

    # Header Row
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for c_idx, h in enumerate(headers):
        cell = tbl.rows[0].cells[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx+1]
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=110, right=110)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def add_code_box(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="14" w:space="0" w:color="0F2942"/>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.1
    r = cp.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

print("Helper definitions loaded.")
