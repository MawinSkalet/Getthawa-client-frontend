import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Human Executive Color Palette (Tailwind Slate / Deep Corporate Navy)
COLOR_TITLE = RGBColor(15, 23, 42)       # Midnight Slate (#0F172A)
COLOR_H1 = RGBColor(15, 23, 42)          # Deep Navy (#0F172A)
COLOR_H2 = RGBColor(30, 58, 138)         # Royal Navy (#1E3A8A)
COLOR_H3 = RGBColor(51, 65, 85)          # Slate (#334155)
COLOR_BODY = RGBColor(30, 41, 59)        # Dark Charcoal Slate (#1E293B)
COLOR_MUTED = RGBColor(100, 116, 139)    # Muted Slate (#64748B)
COLOR_LINK = RGBColor(2, 132, 199)       # Sky Blue (#0284C7)
COLOR_CODE = RGBColor(159, 18, 57)       # Wine Red (#9F1239)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="E2E8F0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_inlines(paragraph, text, base_font_size=10.5, base_color=COLOR_BODY):
    token_pattern = re.compile(
        r'(\*\*[^\*]+?\*\*|\*[^\*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^\)]+?\)|(?:\$\$[\s\S]+?\$\$|\$[^\$]+?\$))'
    )
    parts = token_pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2 and not part.startswith('**'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = COLOR_CODE
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            m = re.match(r'\[([^\]]+?)\]\(([^\)]+?)\)', part)
            if m:
                label, _ = m.groups()
                run = paragraph.add_run(label)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = COLOR_LINK
                run.underline = True
            else:
                run = paragraph.add_run(part)
                run.font.size = Pt(base_font_size)
        elif part.startswith('$') and part.endswith('$'):
            clean_math = part.replace('$$', '').replace('$', '').strip()
            run = paragraph.add_run(clean_math)
            run.font.name = 'Calibri'
            run.italic = True
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = COLOR_H2
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color

def convert_md_to_docx(md_path, docx_path, doc_title):
    doc = docx.Document()
    
    # Standard Executive Page Setup (1-inch Margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Subtle Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"Getthawha Wellness Platform  |  {doc_title}")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        # Subtle Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Commercial Confidential — For Internal & Evaluative Use Only")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_buffer = []

    in_table = False
    table_buffer = []
    
    is_first_h2 = True

    for line_idx, raw_line in enumerate(lines):
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        # Fenced code block check
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_buffer)
                
                # Single-cell elegant code container
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                # Border styling
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:left w:val="single" w:sz="18" w:space="0" w:color="0F172A"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                    </w:tcBorders>
                ''')
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.line_spacing = 1.15
                crun = cp.add_run(code_text)
                crun.font.name = 'Consolas'
                crun.font.size = Pt(8.5)
                crun.font.color.rgb = RGBColor(30, 41, 59)
                
                # Spacing after code block
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)
                
                code_buffer = []
            else:
                if in_table:
                    flush_table(doc, table_buffer)
                    in_table = False
                    table_buffer = []
                in_code_block = True
                code_buffer = []
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        # Table rows check
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_buffer = []
            table_buffer.append(stripped)
            continue
        else:
            if in_table:
                flush_table(doc, table_buffer)
                in_table = False
                table_buffer = []

        # Empty lines
        if not stripped:
            continue

        # Document Title (# )
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = COLOR_TITLE
        # Subtitle or Section Heading (## )
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            if is_first_h2:
                # Subtitle directly under title
                is_first_h2 = False
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(12)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(stripped[3:])
                run.font.name = 'Calibri'
                run.font.size = Pt(13)
                run.font.color.rgb = COLOR_MUTED
            else:
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(stripped[3:])
                run.bold = True
                run.font.name = 'Calibri'
                run.font.size = Pt(14)
                run.font.color.rgb = COLOR_H1
        # Subsection (### )
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_H2
        # Sub-subsection (#### )
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_H3
        elif stripped.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="6" w:color="1E3A8A"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            format_inlines(p, stripped[2:], base_font_size=10.0, base_color=COLOR_MUTED)
        elif stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inlines(p, stripped[2:], base_font_size=10.0)
        elif re.match(r'^\d+\.\s+', stripped):
            m = re.match(r'^\d+\.\s+(.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            format_inlines(p, m.group(1), base_font_size=10.0)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4.5)
            p.paragraph_format.line_spacing = 1.18
            format_inlines(p, stripped, base_font_size=10.5)

    if in_table:
        flush_table(doc, table_buffer)

    doc.save(docx_path)
    print(f"[DOCX] Generated: {docx_path}")

def flush_table(doc, rows_raw):
    parsed_rows = []
    for r in rows_raw:
        cells = [c.strip() for c in r.strip('|').split('|')]
        # Skip divider rows |---|---|
        if all(re.match(r'^:?-+:?$', c) for c in cells if c):
            continue
        parsed_rows.append(cells)

    if not parsed_rows:
        return

    num_rows = len(parsed_rows)
    num_cols = max(len(r) for r in parsed_rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="E2E8F0", sz="4")

    # Repeat header row on every page
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    for r_idx, row_data in enumerate(parsed_rows):
        is_header = (r_idx == 0)
        row = table.rows[r_idx]
        
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.05

            if is_header:
                set_cell_background(cell, "0F172A")
                format_inlines(cp, cell_text, base_font_size=9.0, base_color=RGBColor(255, 255, 255))
                for run in cp.runs:
                    run.bold = True
            else:
                bg = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                format_inlines(cp, cell_text, base_font_size=8.5, base_color=COLOR_BODY)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    docs_dir = os.path.abspath(os.path.join(script_dir, ".."))
    
    files_to_convert = [
        ("1_PROJECT_PROPOSAL.md", "1_PROJECT_PROPOSAL.docx", "Project Proposal & Charter"),
        ("2_SOFTWARE_REQUIREMENTS_SPECIFICATION.md", "2_SOFTWARE_REQUIREMENTS_SPECIFICATION.docx", "Software Requirements Specification"),
        ("3_SOFTWARE_DESIGN_SPECIFICATION.md", "3_SOFTWARE_DESIGN_SPECIFICATION.docx", "Software Design Specification"),
        ("4_SYSTEM_DOCUMENTATION.md", "4_SYSTEM_DOCUMENTATION.docx", "System Documentation & Operations Runbook")
    ]

    for md_name, docx_name, title in files_to_convert:
        md_file = os.path.join(docs_dir, md_name)
        docx_file = os.path.join(docs_dir, docx_name)
        if os.path.exists(md_file):
            convert_md_to_docx(md_file, docx_file, title)
        else:
            print(f"File not found: {md_file}")
