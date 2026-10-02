import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# ==============================================================================
# CORPORATE EXECUTIVE COLOR PALETTE
# ==============================================================================
COLOR_PRIMARY   = RGBColor(15, 41, 66)      # Deep Midnight Navy (#0F2942)
COLOR_SECONDARY = RGBColor(30, 58, 138)     # Royal Navy (#1E3A8A)
COLOR_DARK      = RGBColor(30, 41, 59)      # Charcoal Slate (#1E293B)
COLOR_MUTED     = RGBColor(100, 116, 139)   # Slate Gray (#64748B)
COLOR_BORDER    = "CBD5E1"                  # Slate Border (#CBD5E1)
COLOR_LIGHT_BG  = "F8FAFC"                  # Very Soft Slate (#F8FAFC)
COLOR_CALLOUT_BG= "F1F5F9"                  # Ice Slate Tint (#F1F5F9)
COLOR_CODE      = RGBColor(159, 18, 57)     # Wine Red (#9F1239)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color=COLOR_BORDER, sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_fld_char(run, fld_type):
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), fld_type)
    run._r.append(fld)

def add_instr_text(run, text):
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = text
    run._r.append(instr)

def add_page_number_field(run):
    add_fld_char(run, 'begin')
    add_instr_text(run, 'PAGE')
    add_fld_char(run, 'separate')
    add_fld_char(run, 'end')

def add_total_pages_field(run):
    add_fld_char(run, 'begin')
    add_instr_text(run, 'NUMPAGES')
    add_fld_char(run, 'separate')
    add_fld_char(run, 'end')

def setup_page_layout(doc, doc_title):
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.different_first_page_header_footer = True
    
    # Running Header (Pages 2+)
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(8)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="CBD5E1"/></w:pBdr>')
    hp._p.get_or_add_pPr().append(pBdr)
    
    hrun = hp.add_run(f"Getthawha Wellness Platform  |  {doc_title}")
    hrun.font.name = 'Calibri'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED

    # Running Footer (Pages 2+)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(8)
    pBdr_f = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="4" w:color="CBD5E1"/></w:pBdr>')
    fp._p.get_or_add_pPr().append(pBdr_f)
    
    # Tab stop at 6.5 inches for right alignment
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    
    frun1 = fp.add_run("Commercial Confidential — Getthawha Wellness Group\tPage ")
    frun1.font.name = 'Calibri'
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = COLOR_MUTED
    
    add_page_number_field(frun1)
    
    frun2 = fp.add_run(" of ")
    frun2.font.name = 'Calibri'
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = COLOR_MUTED
    
    add_total_pages_field(frun2)

def create_cover_page(doc, doc_title, subtitle, meta_items):
    # Top institutional badge
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_before = Pt(36)
    p_org.paragraph_format.space_after = Pt(4)
    r_org = p_org.add_run("GETTHAWA THAI MASSAGE & WELLNESS GROUP")
    r_org.font.name = 'Calibri'
    r_org.font.size = Pt(11)
    r_org.font.bold = True
    r_org.font.color.rgb = COLOR_SECONDARY

    # Thin accent rule under org name
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(28)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="14" w:space="1" w:color="0F2942"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.paragraph_format.line_spacing = 1.15
    r_title = p_title.add_run(doc_title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(40)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_MUTED

    # Metadata Panel (Clean 2-column table)
    tbl = doc.add_table(rows=len(meta_items), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(tbl, color="E2E8F0", sz="4")
    
    col_widths = [Inches(2.2), Inches(4.3)]
    for r_idx, (k, v) in enumerate(meta_items):
        row = tbl.rows[r_idx]
        row.cells[0].width = col_widths[0]
        row.cells[1].width = col_widths[1]
        set_cell_margins(row.cells[0], top=70, bottom=70, left=100, right=100)
        set_cell_margins(row.cells[1], top=70, bottom=70, left=100, right=100)
        set_cell_background(row.cells[0], COLOR_LIGHT_BG)
        set_cell_background(row.cells[1], "FFFFFF")

        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK

        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

def create_control_page(doc, revisions, approvals):
    add_heading(doc, "Document Control & Approvals", level=1)
    
    p = doc.add_paragraph("This document constitutes a baseline engineering deliverable for the Getthawha Platform. Formal review, authorization, and historical modifications are recorded below:")
    p.paragraph_format.space_after = Pt(12)

    # Formal Approvals Sign-off Table
    add_heading(doc, "Formal Review & Acceptance Sign-Off", level=2)
    tbl_app = doc.add_table(rows=len(approvals)+1, cols=5)
    tbl_app.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_app)

    headers_app = ["Role / Representation", "Name", "Title / Organization", "Status", "Sign Date"]
    app_col_widths = [Inches(1.5), Inches(1.3), Inches(1.8), Inches(0.9), Inches(1.0)]

    for c_idx, h in enumerate(headers_app):
        cell = tbl_app.rows[0].cells[c_idx]
        cell.width = app_col_widths[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, app in enumerate(approvals):
        row = tbl_app.rows[r_idx+1]
        for c_idx, val in enumerate(app):
            cell = row.cells[c_idx]
            cell.width = app_col_widths[c_idx]
            set_cell_background(cell, COLOR_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(12)

    # Document Revision History Table
    add_heading(doc, "Document Revision History", level=2)
    tbl_rev = doc.add_table(rows=len(revisions)+1, cols=4)
    tbl_rev.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rev)

    headers_rev = ["Version", "Release Date", "Author / Role", "Summary of Technical Changes"]
    rev_col_widths = [Inches(0.9), Inches(1.2), Inches(1.7), Inches(2.7)]

    for c_idx, h in enumerate(headers_rev):
        cell = tbl_rev.rows[0].cells[c_idx]
        cell.width = rev_col_widths[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, rev in enumerate(revisions):
        row = tbl_rev.rows[r_idx+1]
        for c_idx, val in enumerate(rev):
            cell = row.cells[c_idx]
            cell.width = rev_col_widths[c_idx]
            set_cell_background(cell, COLOR_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    sp2 = doc.add_paragraph()
    sp2.paragraph_format.space_after = Pt(14)

    # Confidentiality Notice
    add_heading(doc, "Confidentiality & Proprietary Baseline", level=2)
    p_conf = doc.add_paragraph(
        "The information contained in this document is confidential and proprietary to Getthawha Thai Massage & "
        "Wellness Group. It is submitted with the understanding that it will be held in strict confidence and will "
        "not be disclosed, duplicated, or used, in whole or in part, for any purpose other than evaluating and "
        "operating the Getthawha Omnichannel Platform without prior written consent from the executive board."
    )
    p_conf.paragraph_format.line_spacing = 1.18
    p_conf.paragraph_format.space_after = Pt(6)

    doc.add_page_break()

def create_toc_page(doc, toc_items):
    add_heading(doc, "Table of Contents", level=1)
    p_lead = doc.add_paragraph("This specification is organized into the following formal sections:")
    p_lead.paragraph_format.space_after = Pt(16)

    for num_str, title_str in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        # Tab stop at 0.8 inches for clean alignment
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.8), WD_TAB_ALIGNMENT.LEFT)

        r_num = p.add_run(num_str)
        r_num.font.name = 'Calibri'
        r_num.font.bold = True
        r_num.font.size = Pt(10.5)
        r_num.font.color.rgb = COLOR_PRIMARY

        r_tab = p.add_run("\t")

        r_title = p.add_run(title_str)
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = COLOR_DARK

    doc.add_page_break()

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        # Underline accent
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="0F2942"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_MUTED

def format_inlines(paragraph, text, base_font_size=10.5, base_color=COLOR_DARK):
    token_pattern = re.compile(
        r'(\*\*[^\*]+?\*\*|\*[^\*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^\)]+?\))'
    )
    parts = token_pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2 and not part.startswith('**'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = COLOR_CODE
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            m = re.match(r'\[([^\]]+?)\]\(([^\)]+?)\)', part)
            if m:
                label, _ = m.groups()
                run = paragraph.add_run(label)
                run.font.name = 'Calibri'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = COLOR_SECONDARY
                run.underline = True
            else:
                run = paragraph.add_run(part)
                run.font.size = Pt(base_font_size)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(base_font_size)
            run.font.color.rgb = base_color

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    format_inlines(p, text, base_font_size=10.5, base_color=COLOR_DARK)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    format_inlines(p, text, base_font_size=10, base_color=COLOR_DARK)

def add_numbered(doc, text, num_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    
    r_num = p.add_run(num_str + " ")
    r_num.font.name = 'Calibri'
    r_num.font.bold = True
    r_num.font.size = Pt(10)
    r_num.font.color.rgb = COLOR_PRIMARY

    format_inlines(p, text, base_font_size=10, base_color=COLOR_DARK)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, COLOR_CALLOUT_BG)
    set_cell_margins(cell, top=90, bottom=90, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="0F2942"/>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.15
    
    r_tit = cp.add_run(f"[{title}] ")
    r_tit.font.name = 'Calibri'
    r_tit.font.size = Pt(9.5)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    format_inlines(cp, text, base_font_size=9.5, base_color=COLOR_DARK)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_container(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, COLOR_LIGHT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="16" w:space="0" w:color="0F2942"/>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(0)
    cp.paragraph_format.line_spacing = 1.12
    r = cp.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def render_table(doc, headers, data_rows):
    if not headers or not data_rows:
        return
    num_cols = len(headers)
    tbl = doc.add_table(rows=len(data_rows)+1, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)

    # Header Row XML flags
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    # Calculate smart column widths
    col_max_lens = [len(h) for h in headers]
    for row in data_rows:
        for c_idx in range(min(num_cols, len(row))):
            col_max_lens[c_idx] = max(col_max_lens[c_idx], len(str(row[c_idx])))
    
    # Weight calculation
    raw_weights = [max(l, 5) ** 0.65 for l in col_max_lens]
    total_weight = sum(raw_weights)
    TOTAL_WIDTH = 6.5
    col_widths = [Inches((w / total_weight) * TOTAL_WIDTH) for w in raw_weights]

    for c_idx, h in enumerate(headers):
        cell = tbl.rows[0].cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "0F2942")
        set_cell_margins(cell, top=75, bottom=75, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = 'Calibri'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(data_rows):
        row = tbl.rows[r_idx+1]
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for c_idx in range(num_cols):
            val = str(row_data[c_idx]) if c_idx < len(row_data) else ""
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            bg = COLOR_LIGHT_BG if (r_idx % 2 == 1) else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            format_inlines(p, val, base_font_size=8.5, base_color=COLOR_DARK)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def parse_and_append_markdown_body(doc, md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # Skip top metadata tables / headers until first major section "## 1."
    start_line_idx = 0
    for idx, line in enumerate(lines):
        if line.strip().startswith('## 1.'):
            start_line_idx = idx
            break

    in_code = False
    code_buf = []

    in_table = False
    table_raw_rows = []

    for line in lines[start_line_idx:]:
        stripped = line.strip()

        # Fenced code block
        if stripped.startswith('```'):
            if in_code:
                in_code = False
                code_text = "\n".join(code_buf)
                add_code_container(doc, code_text)
                code_buf = []
            else:
                if in_table:
                    headers, data = process_raw_table(table_raw_rows)
                    render_table(doc, headers, data)
                    in_table = False
                    table_raw_rows = []
                in_code = True
                code_buf = []
            continue

        if in_code:
            code_buf.append(line.rstrip('\r\n'))
            continue

        # Table rows
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_raw_rows = []
            table_raw_rows.append(stripped)
            continue
        else:
            if in_table:
                headers, data = process_raw_table(table_raw_rows)
                render_table(doc, headers, data)
                in_table = False
                table_raw_rows = []

        if not stripped:
            continue

        # Headings
        if stripped.startswith('## '):
            # Major Chapter H1 -> Page Break before each chapter!
            doc.add_page_break()
            add_heading(doc, stripped[3:], level=1)
        elif stripped.startswith('### '):
            add_heading(doc, stripped[4:], level=2)
        elif stripped.startswith('#### '):
            add_heading(doc, stripped[5:], level=3)
        elif stripped.startswith('##### '):
            add_heading(doc, stripped[6:], level=4)
        elif stripped.startswith('> '):
            add_callout(doc, stripped[2:], title="OPERATIONAL CONTEXT")
        elif stripped.startswith('* ') or stripped.startswith('- '):
            add_bullet(doc, stripped[2:])
        elif re.match(r'^\d+\.\s+', stripped):
            m = re.match(r'^(\d+\.)\s+(.*)', stripped)
            if m:
                add_numbered(doc, m.group(2), m.group(1))
            else:
                add_paragraph(doc, stripped)
        else:
            add_paragraph(doc, stripped)

    if in_table:
        headers, data = process_raw_table(table_raw_rows)
        render_table(doc, headers, data)

def process_raw_table(raw_rows):
    cleaned_rows = []
    for r in raw_rows:
        cols = [c.strip() for c in r.strip('|').split('|')]
        # Filter separator row like |:---|:---|
        if any(set(c).issubset({'-', ':', ' '}) for c in cols if c):
            continue
        cleaned_rows.append(cols)
    if not cleaned_rows:
        return [], []
    headers = cleaned_rows[0]
    data = cleaned_rows[1:]
    return headers, data

def build_all_documents():
    docs_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    print(f"Working in docs directory: {docs_dir}")

    # ==============================================================================
    # 1. PROJECT PROPOSAL & CHARTER
    # ==============================================================================
    doc1 = docx.Document()
    setup_page_layout(doc1, "Project Proposal & Charter")
    create_cover_page(
        doc1,
        doc_title="PROJECT PROPOSAL & TECHNICAL CHARTER",
        subtitle="Getthawha Omnichannel Spa & Wellness Booking Platform",
        meta_items=[
            ("Project Title", "Getthawha Omnichannel Spa & Wellness Platform"),
            ("Document Type", "Commercial Proposal & Technical Baseline"),
            ("Standards Alignment", "IEEE Std 1058 / PMBOK 7th Edition Guide"),
            ("Document Version", "Version 2.0 (Enterprise Comprehensive Baseline)"),
            ("Target Organization", "Getthawha Thai Massage & Wellness Group, Chiang Mai"),
            ("Technical Authority", "Enterprise Solutions Architecture Team"),
            ("Effective Date", "September 2026"),
            ("Classification", "Commercial Confidential / Technical Baseline")
        ]
    )
    create_control_page(
        doc1,
        revisions=[
            ["1.0", "August 1, 2026", "Solutions Architect", "Initial feasibility proposal and preliminary architecture"],
            ["1.5", "August 20, 2026", "Product Engineering Lead", "Added multi-channel scope, pricing catalog, and sprint roadmaps"],
            ["2.0", "September 14, 2026", "Chief Software Architect", "Comprehensive revision: Work Breakdown Structure, RACI, Financial ROI, and Risk Register"]
        ],
        approvals=[
            ["Executive Sponsor", "Mawin (Project Owner)", "Founder & Managing Director", "Approved", "September 2026"],
            ["Technical Authority", "Solutions Architect", "Chief Software Architect", "Approved", "September 2026"],
            ["Store Operations", "Operations Director", "Head of Branch Operations", "Approved", "September 2026"],
            ["Quality Assurance", "QA Lead", "Principal Test Architect", "Approved", "September 2026"]
        ]
    )
    create_toc_page(
        doc1,
        toc_items=[
            ("1.0", "Executive Summary & Platform Overview"),
            ("2.0", "Business Case & Problem Statement"),
            ("3.0", "Strategic Objectives & Success Metrics (KPIs)"),
            ("4.0", "Scope Baseline: In-Scope vs. Out-of-Scope"),
            ("5.0", "High-Level Architecture & Technical Approach"),
            ("6.0", "Work Breakdown Structure (WBS) & Implementation Roadmap"),
            ("7.0", "Resource Allocation & RACI Governance Matrix"),
            ("8.0", "Cost-Benefit Analysis & Financial Return on Investment (ROI)"),
            ("9.0", "Risk Management Plan & Mitigation Register"),
            ("10.0", "Acceptance Criteria & Project Sign-Off Baseline")
        ]
    )
    parse_and_append_markdown_body(doc1, os.path.join(docs_dir, "1_PROJECT_PROPOSAL.md"))
    doc1_out = os.path.join(docs_dir, "1_PROJECT_PROPOSAL.docx")
    doc1.save(doc1_out)
    print(f"[OK] Generated {doc1_out}")

    # ==============================================================================
    # 2. SOFTWARE REQUIREMENTS SPECIFICATION (SRS)
    # ==============================================================================
    doc2 = docx.Document()
    setup_page_layout(doc2, "Software Requirements Specification (SRS)")
    create_cover_page(
        doc2,
        doc_title="SOFTWARE REQUIREMENTS SPECIFICATION (SRS)",
        subtitle="Getthawha Omnichannel Spa & Wellness Booking Platform",
        meta_items=[
            ("Project Title", "Getthawha Omnichannel Spa & Wellness Platform"),
            ("Document Type", "Software Requirements Specification (SRS)"),
            ("Standards Compliance", "ISO/IEC/IEEE 29148:2018 / IEEE Std 830-1998"),
            ("Quality Model Alignment", "ISO/IEC 25010 Software Quality Standards"),
            ("Document Version", "Version 2.0 (Comprehensive Baseline)"),
            ("Target Organization", "Getthawha Thai Massage & Wellness Group, Chiang Mai"),
            ("Technical Authority", "Lead Systems Analyst & Solutions Architect"),
            ("Effective Date", "September 2026"),
            ("Classification", "Commercial Confidential / Engineering Specification")
        ]
    )
    create_control_page(
        doc2,
        revisions=[
            ["1.0", "August 10, 2026", "Systems Analyst", "Initial requirements baseline for online booking and admin management"],
            ["1.5", "August 30, 2026", "QA / Product Specialist", "Integrated Google OAuth 2.0 PKCE and LINE LIFF specifications"],
            ["2.0", "September 14, 2026", "Software Engineering Team", "Comprehensive expansion: 15 Functional Requirements, Regex Validation Matrix, and RTM"]
        ],
        approvals=[
            ["Product Owner", "Mawin (Project Owner)", "Executive Stakeholder", "Approved", "September 2026"],
            ["Systems Analysis", "Lead Systems Analyst", "Solutions Architecture", "Approved", "September 2026"],
            ["Engineering", "Lead Software Engineer", "Fullstack Development Team", "Approved", "September 2026"],
            ["Verification", "QA Lead", "Software Quality Assurance", "Approved", "September 2026"]
        ]
    )
    create_toc_page(
        doc2,
        toc_items=[
            ("1.0", "Introduction & Document Conventions"),
            ("2.0", "Overall System Description & Persona Matrix"),
            ("3.0", "Specific Functional Requirements (FR-01 through FR-15)"),
            ("4.0", "Form Input Validation & Business Rule Matrix"),
            ("5.0", "External Interface Requirements (API, GUI, Network)"),
            ("6.0", "Non-Functional Requirements (ISO/IEC 25010 Quality Model)"),
            ("7.0", "Requirements Traceability Matrix (RTM)")
        ]
    )
    parse_and_append_markdown_body(doc2, os.path.join(docs_dir, "2_SOFTWARE_REQUIREMENTS_SPECIFICATION.md"))
    doc2_out = os.path.join(docs_dir, "2_SOFTWARE_REQUIREMENTS_SPECIFICATION.docx")
    doc2.save(doc2_out)
    print(f"[OK] Generated {doc2_out}")

    # ==============================================================================
    # 3. SOFTWARE DESIGN SPECIFICATION (SDS)
    # ==============================================================================
    doc3 = docx.Document()
    setup_page_layout(doc3, "Software Design Specification (SDS)")
    create_cover_page(
        doc3,
        doc_title="SOFTWARE DESIGN SPECIFICATION (SDS)",
        subtitle="Getthawha Omnichannel Spa & Wellness Booking Platform",
        meta_items=[
            ("Project Title", "Getthawha Omnichannel Spa & Wellness Platform"),
            ("Document Type", "Software Design Specification (SDS) / Architecture Baseline"),
            ("Standards Compliance", "IEEE Std 1016-2009 / ISO/IEC/IEEE 42010:2022"),
            ("Architectural Model", "C4 Architectural Viewpoints & Service-Oriented Design"),
            ("Document Version", "Version 2.0 (Comprehensive Baseline)"),
            ("Target Audience", "Software Architects, Fullstack Engineers, DevOps & DBAs"),
            ("Technical Authority", "Enterprise Architecture & Engineering Team"),
            ("Effective Date", "September 2026"),
            ("Classification", "Commercial Confidential / Technical Architecture")
        ]
    )
    create_control_page(
        doc3,
        revisions=[
            ["1.0", "August 15, 2026", "Solutions Architect", "Initial design baseline: container topology, ERD, and core sequence"],
            ["1.5", "September 2, 2026", "Backend Lead", "Integrated Google OAuth 2.0 PKCE and Meta Facebook Webhook architecture"],
            ["2.0", "September 14, 2026", "Architecture Team", "Full enterprise expansion: 26 API data contracts, database index plan, and concurrency locks"]
        ],
        approvals=[
            ["Chief Architect", "Lead Software Architect", "Enterprise Solutions", "Approved", "September 2026"],
            ["Backend Engineering", "Senior Backend Engineer", "API Platform Team", "Approved", "September 2026"],
            ["Frontend Engineering", "Senior Frontend Engineer", "Client/Admin Web Team", "Approved", "September 2026"],
            ["Infrastructure", "DevOps & Cloud Lead", "Site Reliability Engineering", "Approved", "September 2026"]
        ]
    )
    create_toc_page(
        doc3,
        toc_items=[
            ("1.0", "Architectural Principles & Design Goals"),
            ("2.0", "System Architecture & C4 Viewpoints (L1 to L3)"),
            ("3.0", "Relational Data Model & PostgreSQL Data Dictionary"),
            ("4.0", "Concurrency Control & ACID Transaction Design"),
            ("5.0", "REST API Data Contracts & Endpoint Specifications (26 Endpoints)"),
            ("6.0", "Security Architecture & Authentication Flows (OAuth 2.0 PKCE & JWT)"),
            ("7.0", "Third-Party Integration Architecture (LINE LIFF & Meta Graph)"),
            ("8.0", "Error Handling, Logging, and Observability Architecture")
        ]
    )
    parse_and_append_markdown_body(doc3, os.path.join(docs_dir, "3_SOFTWARE_DESIGN_SPECIFICATION.md"))
    doc3_out = os.path.join(docs_dir, "3_SOFTWARE_DESIGN_SPECIFICATION.docx")
    doc3.save(doc3_out)
    print(f"[OK] Generated {doc3_out}")

    # ==============================================================================
    # 4. SYSTEM DOCUMENTATION & OPERATIONS RUNBOOK
    # ==============================================================================
    doc4 = docx.Document()
    setup_page_layout(doc4, "System Documentation & Operations Runbook")
    create_cover_page(
        doc4,
        doc_title="SYSTEM DOCUMENTATION & OPERATIONS RUNBOOK",
        subtitle="Getthawha Omnichannel Spa & Wellness Booking Platform",
        meta_items=[
            ("Project Title", "Getthawha Omnichannel Spa & Wellness Platform"),
            ("Document Type", "System Documentation & SRE Operations Runbook"),
            ("Standards Compliance", "ISO/IEC 26514:2008 / ISO/IEC 27001 (SecOps & Runbooks)"),
            ("Document Version", "Version 2.0 (Comprehensive Baseline)"),
            ("Target Audience", "DevOps Engineers, SREs, System Administrators, Developers"),
            ("Technical Authority", "DevOps & Site Reliability Engineering Team"),
            ("Effective Date", "September 2026"),
            ("Classification", "Commercial Confidential / Operational Baseline")
        ]
    )
    create_control_page(
        doc4,
        revisions=[
            ["1.0", "August 20, 2026", "Infrastructure Lead", "Baseline Docker Compose deployment and environment tables"],
            ["1.5", "September 5, 2026", "SRE Specialist", "Integrated Cloudflare Tunnel setup and Facebook Bot verification scripts"],
            ["2.0", "September 14, 2026", "DevOps & Operations Team", "Full enterprise expansion: Click-by-click console SOPs, disaster recovery, and 8 incident runbooks"]
        ],
        approvals=[
            ["Operations Director", "Mawin (Project Owner)", "Executive Operations", "Approved", "September 2026"],
            ["SRE Lead", "Lead DevOps Engineer", "Infrastructure Operations", "Approved", "September 2026"],
            ["Security Officer", "InfoSec Lead", "Security & Compliance", "Approved", "September 2026"],
            ["Release Manager", "Product Delivery Lead", "Release Engineering", "Approved", "September 2026"]
        ]
    )
    create_toc_page(
        doc4,
        toc_items=[
            ("1.0", "System Overview & Multi-Repository Monorepo Structure"),
            ("2.0", "Comprehensive Environment Configuration Dictionary (.env)"),
            ("3.0", "Deployment Guide (Docker Compose & Cloud Environments)"),
            ("4.0", "Third-Party Console Setup SOPs (Google Cloud, LINE Developers, Meta Developers)"),
            ("5.0", "Database Administration & Disaster Recovery Runbook"),
            ("6.0", "Site Reliability & Production Incident Runbooks (IR-01 through IR-08)")
        ]
    )
    parse_and_append_markdown_body(doc4, os.path.join(docs_dir, "4_SYSTEM_DOCUMENTATION.md"))
    doc4_out = os.path.join(docs_dir, "4_SYSTEM_DOCUMENTATION.docx")
    doc4.save(doc4_out)
    print(f"[OK] Generated {doc4_out}")

if __name__ == "__main__":
    build_all_documents()
